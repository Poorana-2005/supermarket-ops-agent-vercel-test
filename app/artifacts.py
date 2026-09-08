from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from pptx import Presentation
from pptx.util import Inches
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from .config import ARTIFACT_DIR, SHOP_NAME, SHOP_GSTIN
FONT_PATH = matplotlib.get_data_path() + "/fonts/ttf/DejaVuSans.ttf"
BOLD_FONT_PATH = matplotlib.get_data_path() + "/fonts/ttf/DejaVuSans-Bold.ttf"

pdfmetrics.registerFont(TTFont("DejaVuSans", FONT_PATH))
pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", BOLD_FONT_PATH))

def create_invoice(sale):
    path=ARTIFACT_DIR / f"{sale['bill_no']}.pdf"
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=28,leftMargin=28,topMargin=28,bottomMargin=28)
    styles=getSampleStyleSheet()
    styles["Normal"].fontName = "DejaVuSans"
    styles["Title"].fontName = "DejaVuSans-Bold"
    story=[Paragraph(f"<b>{SHOP_NAME}</b>",styles["Title"]),
           Paragraph(f"GSTIN: {SHOP_GSTIN}",styles["Normal"]),
           Paragraph(f"Tax Invoice: {sale['bill_no']} | {sale['created_at']}",styles["Normal"]),
           Spacer(1,12)]
    data=[["Item","HSN","Qty","Rate","Taxable","CGST","SGST","Total"]]
    for x in sale["items"]:
        data.append([x["name"],x["hsn"],x["qty"],f"₹{x['rate']:.2f}",f"₹{x['taxable_value']:.2f}",
                     f"₹{x['cgst']:.2f}",f"₹{x['sgst']:.2f}",f"₹{x['line_total']:.2f}"])
    data += [["","","","Subtotal",f"₹{sale['subtotal']:.2f}","","",""],
             ["","","","Tax",f"₹{sale['total_tax']:.2f}","","",""],
             ["","","","Grand Total",f"₹{sale['grand_total']:.2f}","","",""]]
    table=Table(data,repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("GRID",(0,0),(-1,-1),0.4,colors.grey),
        ("FONTNAME",(0,0),(-1,0),"DejaVuSans-Bold"),
        ("FONTNAME",(0,1),(-1,-1),"DejaVuSans"),("ALIGN",(2,1),(-1,-1),"RIGHT"),
        ("FONTNAME",(3,-3),(-1,-1),"DejaVuSans-Bold"),("BOTTOMPADDING",(0,0),(-1,0),7)
    ]))
    story += [table,Spacer(1,12),
              Paragraph(f"Payment: {sale['payment_mode']} | Reference: {sale['payment_reference']}",styles["Normal"]),
              Paragraph(f"CGST: ₹{sale['cgst']:.2f} | SGST: ₹{sale['sgst']:.2f} | Total GST: ₹{sale['total_tax']:.2f}",styles["Normal"])]
    doc.build(story)
    return str(path)

def create_analysis_deck(summary):
    chart=ARTIFACT_DIR/"weekly_sales.png"
    plt.figure(figsize=(8,4.5))
    labels=[x["day"] for x in summary["daily"]]
    values=[x["revenue"] for x in summary["daily"]]
    plt.bar(labels, values)
    plt.title("7-Day Revenue Trend")
    plt.xlabel("Date")
    plt.ylabel("Revenue (₹)")
    plt.tight_layout()
    plt.savefig(chart,dpi=160)
    plt.close()

    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text=f"{SHOP_NAME} — Ops Analysis"
    slide.placeholders[1].text="Sales, tax, payment mix and stock health"

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text="Revenue"
    slide.shapes.add_picture(str(chart), Inches(0.8), Inches(1.3), width=Inches(8.5))

    slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text="Top Products"
    tx=slide.shapes.add_textbox(Inches(0.8),Inches(1.4),Inches(8),Inches(4)).text_frame
    for i,x in enumerate(summary["top_items"][:5]):
        p=tx.paragraphs[0] if i==0 else tx.add_paragraph()
        p.text=f"{i+1}. {x['name']} — {x['qty']} units"
    slide=prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text="Store Health"
    tf=slide.shapes.add_textbox(Inches(0.8),Inches(1.4),Inches(8),Inches(4)).text_frame
    insights=[
        f"Revenue: ₹{summary['total']:.2f}",
        f"GST collected: ₹{summary['tax']:.2f}",
        f"Cash: ₹{summary['payments'].get('CASH',0):.2f}",
        f"UPI: ₹{summary['payments'].get('UPI',0):.2f}",
        f"Card: ₹{summary['payments'].get('CARD',0):.2f}",
        f"Low-stock SKUs: {summary['low_stock_count']}",
    ]
    for i,t in enumerate(insights):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=t
    path = ARTIFACT_DIR / f"kirana_ops_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
    prs.save(path)
    return str(path)
